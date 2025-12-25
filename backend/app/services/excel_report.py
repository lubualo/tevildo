"""
Excel report generation service.

Generates accountant-ready Excel reports with invoice summaries.
"""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from app.models.invoice import Invoice
from app.models.settings import UserSettings


class ExcelReportGenerator:
    """Generates Excel reports for Monotributo tracking."""

    def __init__(
        self,
        invoices: List[Invoice],
        settings: UserSettings,
        date_from: date,
        date_to: date,
        category_info: Optional[dict] = None,
    ):
        self.invoices = invoices
        self.settings = settings
        self.date_from = date_from
        self.date_to = date_to
        self.category_info = category_info
        self.wb = Workbook()

        # Styles
        self.header_font = Font(bold=True, size=12)
        self.title_font = Font(bold=True, size=14)
        self.currency_format = '#,##0.00'
        self.date_format = 'DD/MM/YYYY'
        self.header_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def generate(self) -> BytesIO:
        """Generate the complete Excel report."""
        # Remove default sheet
        self.wb.remove(self.wb.active)

        # Create sheets
        self._create_summary_sheet()
        self._create_invoices_sheet()
        self._create_monthly_sheet()

        # Save to BytesIO
        output = BytesIO()
        self.wb.save(output)
        output.seek(0)
        return output

    def _create_summary_sheet(self):
        """Create the summary sheet with key metrics."""
        ws = self.wb.create_sheet("Summary")

        # Title
        ws['A1'] = "Monotributo Income Report"
        ws['A1'].font = Font(bold=True, size=16)
        ws.merge_cells('A1:D1')

        # Company info
        row = 3
        ws[f'A{row}'] = "Company:"
        ws[f'A{row}'].font = self.header_font
        ws[f'B{row}'] = self.settings.company_name or "Not set"

        row += 1
        ws[f'A{row}'] = "Income Type:"
        ws[f'A{row}'].font = self.header_font
        ws[f'B{row}'] = "Services" if self.settings.income_type == "services" else "Goods"

        row += 1
        ws[f'A{row}'] = "Current Category:"
        ws[f'A{row}'].font = self.header_font
        ws[f'B{row}'] = self.settings.current_category or "Not set"

        # Report period
        row += 2
        ws[f'A{row}'] = "Report Period"
        ws[f'A{row}'].font = self.title_font
        ws.merge_cells(f'A{row}:D{row}')

        row += 1
        ws[f'A{row}'] = "From:"
        ws[f'A{row}'].font = self.header_font
        ws[f'B{row}'] = self.date_from.strftime("%d/%m/%Y")

        row += 1
        ws[f'A{row}'] = "To:"
        ws[f'A{row}'].font = self.header_font
        ws[f'B{row}'] = self.date_to.strftime("%d/%m/%Y")

        # Totals
        row += 2
        ws[f'A{row}'] = "Summary"
        ws[f'A{row}'].font = self.title_font
        ws.merge_cells(f'A{row}:D{row}')

        total_ars = sum(float(inv.total_ars or 0) for inv in self.invoices)

        row += 1
        ws[f'A{row}'] = "Total Invoiced (ARS):"
        ws[f'A{row}'].font = self.header_font
        ws[f'B{row}'] = total_ars
        ws[f'B{row}'].number_format = self.currency_format

        row += 1
        ws[f'A{row}'] = "Invoice Count:"
        ws[f'A{row}'].font = self.header_font
        ws[f'B{row}'] = len(self.invoices)

        # Category usage if available
        if self.category_info and self.settings.current_category:
            row += 2
            ws[f'A{row}'] = "Category Usage"
            ws[f'A{row}'].font = self.title_font
            ws.merge_cells(f'A{row}:D{row}')

            row += 1
            ws[f'A{row}'] = "Category Limit:"
            ws[f'A{row}'].font = self.header_font
            ws[f'B{row}'] = self.category_info.get('limit', 0)
            ws[f'B{row}'].number_format = self.currency_format

            row += 1
            ws[f'A{row}'] = "Usage:"
            ws[f'A{row}'].font = self.header_font
            ws[f'B{row}'] = f"{self.category_info.get('percentage', 0):.1f}%"

            row += 1
            ws[f'A{row}'] = "Remaining:"
            ws[f'A{row}'].font = self.header_font
            ws[f'B{row}'] = self.category_info.get('remaining', 0)
            ws[f'B{row}'].number_format = self.currency_format

        # Adjust column widths
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 20

    def _create_invoices_sheet(self):
        """Create the detailed invoices sheet."""
        ws = self.wb.create_sheet("Invoices")

        # Headers
        headers = [
            "Issue Date",
            "Invoice #",
            "Type",
            "Client",
            "Client Tax ID",
            "Currency",
            "Original Amount",
            "FX Rate",
            "Amount (ARS)",
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.border = self.thin_border
            cell.alignment = Alignment(horizontal='center')

        # Data rows
        for row_num, inv in enumerate(self.invoices, 2):
            ws.cell(row=row_num, column=1, value=inv.issue_date)
            ws.cell(row=row_num, column=2, value=inv.invoice_number or "")
            ws.cell(row=row_num, column=3, value=inv.invoice_type or "")
            ws.cell(row=row_num, column=4, value=inv.client_name or "")
            ws.cell(row=row_num, column=5, value=inv.client_tax_id or "")
            ws.cell(row=row_num, column=6, value=inv.currency or "ARS")

            orig_cell = ws.cell(row=row_num, column=7, value=float(inv.total_original or 0))
            orig_cell.number_format = self.currency_format

            fx_cell = ws.cell(row=row_num, column=8, value=float(inv.fx_rate or 1))
            fx_cell.number_format = '#,##0.0000'

            ars_cell = ws.cell(row=row_num, column=9, value=float(inv.total_ars or 0))
            ars_cell.number_format = self.currency_format

            # Apply borders
            for col in range(1, 10):
                ws.cell(row=row_num, column=col).border = self.thin_border

        # Adjust column widths
        widths = [12, 20, 8, 30, 15, 10, 18, 12, 18]
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width

        # Add total row
        if self.invoices:
            total_row = len(self.invoices) + 2
            ws.cell(row=total_row, column=8, value="TOTAL:").font = self.header_font
            total_cell = ws.cell(
                row=total_row,
                column=9,
                value=sum(float(inv.total_ars or 0) for inv in self.invoices)
            )
            total_cell.font = self.header_font
            total_cell.number_format = self.currency_format

    def _create_monthly_sheet(self):
        """Create monthly breakdown sheet."""
        ws = self.wb.create_sheet("Monthly Breakdown")

        # Group invoices by month
        monthly_totals = {}
        for inv in self.invoices:
            if inv.issue_date:
                # Parse YYYY-MM-DD format
                year_month = inv.issue_date[:7]  # "YYYY-MM"
                if year_month not in monthly_totals:
                    monthly_totals[year_month] = {"count": 0, "total": 0}
                monthly_totals[year_month]["count"] += 1
                monthly_totals[year_month]["total"] += float(inv.total_ars or 0)

        # Sort by month
        sorted_months = sorted(monthly_totals.keys())

        # Headers
        headers = ["Month", "Invoice Count", "Total (ARS)", "Cumulative Total"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.border = self.thin_border
            cell.alignment = Alignment(horizontal='center')

        # Data rows
        cumulative = 0
        for row_num, month in enumerate(sorted_months, 2):
            data = monthly_totals[month]
            cumulative += data["total"]

            # Format month as "Month YYYY"
            year, month_num = month.split("-")
            month_names = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
                          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
            month_display = f"{month_names[int(month_num)-1]} {year}"

            ws.cell(row=row_num, column=1, value=month_display)
            ws.cell(row=row_num, column=2, value=data["count"])

            total_cell = ws.cell(row=row_num, column=3, value=data["total"])
            total_cell.number_format = self.currency_format

            cum_cell = ws.cell(row=row_num, column=4, value=cumulative)
            cum_cell.number_format = self.currency_format

            # Apply borders
            for col in range(1, 5):
                ws.cell(row=row_num, column=col).border = self.thin_border

        # Adjust column widths
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 18
